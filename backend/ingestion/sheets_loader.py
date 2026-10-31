"""Load raw sheet cells verbatim (strings) plus col-A hyperlink and sheet row number."""
import hashlib
import openpyxl
from backend import config

COLS = ["docket", "subject", "status", "cfr_part", "issue_date", "service_office", "office_primary",
        "ad_number", "publish_date", "comments_due", "make", "model", "product_type", "product_subtype",
        "ac_ref", "air_policy_ref", "additional_ref", "col_r", "citation", "page_number", "summary", "comments"]


def verify_sha():
    h = hashlib.sha256(config.SHEET_XLSX.read_bytes()).hexdigest()
    if h != config.SHEET_SHA256:
        raise RuntimeError(f"dataset.xlsx sha256 mismatch: {h}")
    return h


def load_rows():
    verify_sha()
    ws = openpyxl.load_workbook(config.SHEET_XLSX).active
    rows = []
    for r in range(2, ws.max_row + 1):
        d = {"row": r}
        for i, name in enumerate(COLS, 1):
            v = ws.cell(r, i).value
            d[name] = "" if v is None else str(v)
        hl = ws.cell(r, 1).hyperlink
        d["drs_url"] = hl.target if hl and hl.target else ""
        rows.append(d)
    return rows
