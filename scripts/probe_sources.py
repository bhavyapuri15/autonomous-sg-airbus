#!/usr/bin/env python3
"""Phase 0.4 source probe: FR metadata for all sheet rows, full-text sample, Drive listing.
Idempotent and cache-aware: cached files are never refetched. Prints network-call count."""
import csv, hashlib, json, re, sys, time, datetime as dt
from pathlib import Path
import openpyxl, requests

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data/raw/sheet_snapshot/dataset.xlsx"
META = ROOT / "data/raw/fr_metadata"
DOCS = ROOT / "data/raw/fr_docs"
OUT = ROOT / "data/processed/source_manifest.csv"
API = "https://www.federalregister.gov/api/v1"
UA = {"User-Agent": "autonomous-sg-airbus-hackathon-research-prototype/0.1 (read-only FAA AD metadata check)"}
NET = {"calls": 0}
URLRE = re.compile(r"federalregister\.gov/(?:documents/\d{4}/\d\d/\d\d|d|a)/((?:[A-Z]\d?-)?\d{2,4}-\d+)")
SAMPLE_N = 5

def get(url, params=None, retries=4):
    for i in range(retries):
        NET["calls"] += 1
        try:
            r = requests.get(url, params=params, headers=UA, timeout=60)
            if r.status_code in (429, 500, 502, 503, 504):
                time.sleep(2 ** (i + 1)); continue
            time.sleep(0.4)
            return r
        except requests.RequestException as e:
            time.sleep(2 ** (i + 1)); last = e
    return None

MONTHS = "January|February|March|April|May|June|July|August|September|October|November|December"
def parse_date(s):
    m = re.search(r"(%s) (\d{1,2}), (\d{4})" % MONTHS, s or "")
    return dt.datetime.strptime(m.group(0), "%B %d, %Y").date().isoformat() if m else None

DRSRE = re.compile(r"FR-ADNPRM-((?:[A-Z]\d?-)?\d{4}-\d+)-")
def read_rows():
    ws = openpyxl.load_workbook(XLSX, read_only=False).active
    rows = []
    for r in range(2, ws.max_row + 1):
        v = lambda c: ("" if ws.cell(r, c).value is None else str(ws.cell(r, c).value))
        link = v(18)
        hl = ws.cell(r, 1).hyperlink.target if ws.cell(r, 1).hyperlink else ""
        mr = URLRE.search(link)  # first FR URL in col R (documents/, /d/, /a/)
        md = DRSRE.search(hl or "")
        if mr:
            kind = "colR_documents_url" if "/documents/" in link[:mr.end()] else ("colR_d_url" if "/d/" in link[:mr.end()] else "colR_a_url")
            docnum = mr.group(1)
        elif md:
            docnum, kind = md.group(1), "colA_drs_link"
        else:
            docnum, kind = "", ""
        pages = re.search(r"(\d+)\s*-\s*(\d+)", v(20))
        d = re.search(r"FAA-\d{4}-\d+", v(1))
        rows.append(dict(row=r, docket=d.group(0) if d else v(1)[:40], link=hl, docnum=docnum, docnum_source=kind,
                         extra_docnums=sorted(set(URLRE.findall(link)) - {docnum}),
                         sheet_date=parse_date(v(19)) or None, sheet_cite_date_col=v(9),
                         sheet_pages=(int(pages.group(1)), int(pages.group(2))) if pages else None))
    return rows

FIELDS = ["document_number", "start_page", "end_page", "title", "type", "publication_date"]
SEARCH = ROOT / "data/raw/fr_search"
def _search(key, params):
    p = SEARCH / (re.sub(r"[^\w.-]", "_", key) + ".json")
    if p.exists():
        return json.loads(p.read_text())
    r = get(f"{API}/documents.json", params=params)
    if r is None or r.status_code != 200:
        return None
    data = r.json(); p.write_text(json.dumps(data)); return data

def resolve_by_search(row):
    """Docket + publication date (+-2 days) search. Returns (docnum, status). Multi-hit not disambiguated by start page -> AMBIGUOUS."""
    if not row["sheet_date"]:
        return "", "no_url;no_sheet_date"
    d0 = dt.date.fromisoformat(row["sheet_date"])
    lo, hi = (d0 - dt.timedelta(days=2)).isoformat(), (d0 + dt.timedelta(days=2)).isoformat()
    dk = row["docket"]
    base = {"conditions[publication_date][gte]": lo, "conditions[publication_date][lte]": hi, "per_page": 20, "fields[]": FIELDS}
    for how, extra in (("docket_id", {"conditions[docket_id]": dk}),
                       ("term", {"conditions[term]": dk, "conditions[agencies][]": "federal-aviation-administration"})):
        data = _search(f"{how}_{dk}_{lo}_{hi}", {**base, **extra})
        if data is None:
            return "", f"search_error_{how}"
        res = data.get("results", [])
        if not res:
            continue
        if len(res) > 1 and row["sheet_pages"]:
            pm = [x for x in res if x.get("start_page") == row["sheet_pages"][0]]
            if len(pm) == 1: res = pm
        if len(res) == 1:
            return res[0]["document_number"], f"search_{how}"
        return "", f"AMBIGUOUS_{how}({len(res)} hits:{'|'.join(x['document_number'] for x in res)})"
    return "", "search_no_match"

def fetch_meta(docnums):
    todo = [d for d in docnums if not (META / f"{d}.json").exists()]
    status = {}
    for i in range(0, len(todo), 20):
        batch = todo[i:i + 20]
        r = get(f"{API}/documents/{','.join(batch)}.json")
        got = set()
        if r is not None and r.status_code == 200:
            for doc in r.json().get("results", []):
                (META / f"{doc['document_number']}.json").write_text(json.dumps(doc)); got.add(doc["document_number"])
        for d in batch:  # individual retry for anything the batch did not return
            if d not in got:
                r2 = get(f"{API}/documents/{d}.json")
                if r2 is not None and r2.status_code == 200:
                    (META / f"{d}.json").write_text(r2.text)
                else:
                    status[d] = f"http_{getattr(r2, 'status_code', 'neterr')}"
        print(f"  meta batch {i//20+1}/{(len(todo)+19)//20} done", file=sys.stderr)
    return status

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()

def fetch_fulltext(docnum, meta):
    d = DOCS / docnum; d.mkdir(parents=True, exist_ok=True)
    out = {}
    for key, url, name in [("xml", meta.get("full_text_xml_url"), "full_text.xml"),
                           ("pdf", meta.get("pdf_url"), "document.pdf"),
                           ("txt", meta.get("raw_text_url"), "raw.txt")]:
        p = d / name
        if url and not p.exists():
            r = get(url)
            if r is not None and r.status_code == 200:
                p.write_bytes(r.content)
            else:
                out[key] = f"FAILED http_{getattr(r, 'status_code', 'neterr')}"; continue
        if p.exists():
            out[key] = dict(file=str(p.relative_to(ROOT)), bytes=p.stat().st_size, sha256=sha(p), url=url)
    return out

def drive_listing():
    p = ROOT / "data/raw/drive_listing.html"
    if not p.exists():
        r = get("https://drive.google.com/embeddedfolderview?id=13h0JZnSqhYTX3rV8dGDGy5AiwaADQvJR")
        if r is None: return {"status": "neterr"}
        p.write_text(r.text); (ROOT / "data/raw/drive_listing.status").write_text(str(r.status_code))
    html = p.read_text()
    items = re.findall(r'id="entry-([\w-]+)".*?flip-entry-title">([^<]+)<', html, re.S)
    return {"status": (ROOT / "data/raw/drive_listing.status").read_text(), "items": items, "bytes": len(html)}

def main():
    META.mkdir(parents=True, exist_ok=True); SEARCH.mkdir(parents=True, exist_ok=True); DOCS.mkdir(parents=True, exist_ok=True); OUT.parent.mkdir(parents=True, exist_ok=True)
    rows = read_rows()
    for r in rows:
        r["resolution"] = r["docnum_source"]
        if not r["docnum"]:
            r["docnum"], r["resolution"] = resolve_by_search(r); r["search_note"] = r["resolution"]
    uniq = sorted({r["docnum"] for r in rows if r["docnum"]})
    print(f"rows={len(rows)} unique_docnums={len(uniq)}", file=sys.stderr)
    status = fetch_meta(uniq)
    metas, summary = {}, {}
    for r in rows:
        n = r["docnum"]; notes = []
        notes.append("docnum_source=" + (r["resolution"] or "none"))
        if r["extra_docnums"]: notes.append("extra_docnums_in_cell=" + "|".join(r["extra_docnums"]))
        m = json.loads((META / f"{n}.json").read_text()) if n and (META / f"{n}.json").exists() else None
        if m:
            metas[n] = m; fs = "ok"
            if not any(r["docket"] in x for x in (m.get("docket_ids") or [])): notes.append(f"docket_mismatch(sheet={r['docket']},fr={m.get('docket_ids')})")
            if r["sheet_date"] and r["sheet_date"] != m.get("publication_date"): notes.append(f"date_mismatch(sheet={r['sheet_date']},fr={m.get('publication_date')})")
            if r["sheet_pages"] and r["sheet_pages"] != (m.get("start_page"), m.get("end_page")): notes.append(f"pages_mismatch(sheet={r['sheet_pages']},fr={(m.get('start_page'), m.get('end_page'))})")
            if not r["sheet_date"]: notes.append("sheet_date_unparsed")
        else:
            fs = "no_docnum" if not n else "failed_" + status.get(n, "unknown")
        r.update(fs=fs, m=m or {}, notes="; ".join(notes))
    with open(OUT, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["row", "docket", "docnum", "url", "fr_fetch_status", "fr_type", "start_page", "end_page", "page_length", "pdf_url", "xml_url", "docnum_source", "metadata_mismatch_notes"])
        for r in rows:
            m = r["m"]
            w.writerow([r["row"], r["docket"], r["docnum"], r["link"].replace("\n", " ")[:200], r["fs"], m.get("type", ""), m.get("start_page", ""), m.get("end_page", ""),
                        m.get("page_length", ""), m.get("pdf_url", ""), m.get("full_text_xml_url", ""), r["resolution"], r["notes"]])
    # sample of 5 unique docs spread across the dataset (from resolved ok docs, in sheet order)
    ok = [r["docnum"] for r in rows if r["fs"] == "ok"]; ok = list(dict.fromkeys(ok))
    sample = [ok[round(i * (len(ok) - 1) / (SAMPLE_N - 1))] for i in range(SAMPLE_N)] if ok else []
    ft = {n: fetch_fulltext(n, metas[n]) for n in sample}
    (ROOT / "data/processed/fulltext_sample_hashes.json").write_text(json.dumps(ft, indent=1))
    dl = drive_listing()
    # stats
    from collections import Counter
    print("fetch_status", dict(Counter(r["fs"] for r in rows)))
    print("resolution", dict(Counter(re.sub(r"\(.*","",r["resolution"]) for r in rows)))
    print("types(unique docs)", dict(Counter(m.get("type") for m in metas.values())))
    print("sum page_length unique docs", sum(m.get("page_length") or 0 for m in metas.values()), "n_docs", len(metas))
    mm = Counter(k.split("(")[0] for r in rows for k in r["notes"].split("; ") if "mismatch" in k)
    print("mismatch counts", dict(mm))
    print("sample", sample); print("drive", dl.get("status"), dl.get("items"))
    print(f"NETWORK_CALLS={NET['calls']}")

if __name__ == "__main__":
    main()
