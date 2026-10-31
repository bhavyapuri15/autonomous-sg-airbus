#!/usr/bin/env python3
"""Profile the AD NPRM spreadsheet. Usage: inspect_dataset.py <xlsx> [out.json]"""
import sys, re, json, statistics
from collections import Counter
from datetime import datetime
import openpyxl

path = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else "data/processed/schema_profile.json"
wb = openpyxl.load_workbook(path)  # not read_only: need hyperlinks
ws = wb.worksheets[0]
hdr = [c.value for c in ws[1]]
rows, links = [], []
for r in ws.iter_rows(min_row=2):
    rows.append([("" if c.value is None else str(c.value)) for c in r])
    links.append({h: c.hyperlink.target for h, c in zip(hdr, r) if c.hyperlink})
N = len(rows)
col = {h: [r[i] for r in rows] for i, h in enumerate(hdr)}
emp = lambda v: v.strip() == ""
tr = lambda s, n=100: s if len(s) <= n else s[:n] + "..."
P = {"sheet": ws.title, "dims": ws.dimensions, "n_rows": N, "headers": hdr}

DATECOLS = ["Issue Date", "Citation Publish Date", "Comments Due Date"]
def pdate(s):
    try: return datetime.strptime(s.strip(), "%m/%d/%Y")
    except Exception: return None

cols = {}
for h in hdr:
    v = col[h]; ne = [x for x in v if not emp(x)]
    d = {"non_empty": len(ne), "empty": N - len(ne), "unique": len(set(ne)),
         "top5": [[tr(k), c] for k, c in sorted(Counter(ne).items(), key=lambda kv: (-kv[1], kv[0]))[:5]]}
    if ne:
        L = [len(x) for x in ne]; d["len_min"], d["len_max"] = min(L), max(L)
    if h in DATECOLS:
        ds = [pdate(x) for x in ne]; ok = [x for x in ds if x]
        d["type"] = "date"; d["date_format"] = "%m/%d/%Y"
        d["parse_ok"] = len(ok); d["parse_fail"] = len(ne) - len(ok)
        if ok: d["min"], d["max"] = min(ok).strftime("%Y-%m-%d"), max(ok).strftime("%Y-%m-%d")
    else:
        d["type"] = "text"
    cols[h] = d
P["columns"] = cols

# Hyperlinks
hl = Counter(k for l in links for k in l)
P["hyperlinks"] = {"cells_by_column": dict(hl), "rows_with_hyperlink": sum(1 for l in links if l)}
tg = [l.get("Docket Number", "") for l in links]
P["hyperlinks"]["docket_target_unique"] = len(set(t for t in tg if t))
P["hyperlinks"]["docket_target_samples"] = [t for t in tg[:3]]
P["hyperlinks"]["docket_target_pattern_ok"] = sum(bool(re.fullmatch(r"https://drs\.faa\.gov/browse/excelExternalWindow/FR-ADNPRM-\d{4}-\d+-\d+\.\d+", t)) for t in tg)
P["hyperlinks"]["docket_target_FR_docnum_ok"] = sum(bool(re.search(r"FR-ADNPRM-(\d{4}-\d+)-", t)) for t in tg)
urls = col["Link for Submitting Public Comment"]
LR = [len(u) for u in urls if u]
P["link_text_column"] = {"len_min": min(LR), "len_max": max(LR), "max_len_count": sum(l == max(LR) for l in LR),
    "starts_https_federalregister": sum(u.startswith("https://www.federalregister.gov/documents/") for u in urls),
    "has_hyperlink_object": hl.get("Link for Submitting Public Comment", 0),
    "len_hist": dict(sorted(Counter(min(l // 20 * 20, 200) for l in LR).items()))}
pat = re.compile(r"/documents/(\d{4})/(\d{2})/(\d{2})/([^/]+)/")
fr = []
for u in urls:
    m = pat.search(u); fr.append(m.group(4) if m else None)
frs = [x for x in fr if x]
P["fr_docnum"] = {"rows_extracted": len(frs), "unique": len(set(frs)), "samples": frs[:3],
    "format_ok_YYYY-NNNNN": sum(bool(re.fullmatch(r"\d{4}-\d{4,6}", x)) for x in frs)}
# docnum from drs target vs FR URL
mism = 0; cmp_n = 0
for t, f in zip(tg, fr):
    m = re.search(r"FR-ADNPRM-(\d{4}-\d+)-", t or "")
    if m and f:
        cmp_n += 1; mism += m.group(1) != f
P["fr_docnum"]["drs_vs_url_compared"] = cmp_n; P["fr_docnum"]["drs_vs_url_mismatch"] = mism
# URL date vs Citation Publish Date
dm = 0; dc = 0
for u, cp in zip(urls, col["Citation Publish Date"]):
    m = pat.search(u); d = pdate(cp)
    if m and d:
        dc += 1; dm += (f"{m.group(1)}-{m.group(2)}-{m.group(3)}" != d.strftime("%Y-%m-%d"))
P["url_date_vs_publish_date"] = {"compared": dc, "mismatch": dm}

# Duplicates
def dups(vals):
    c = Counter(v for v in vals if not emp(v)); d = {k: n for k, n in c.items() if n > 1}
    return {"dup_values": len(d), "rows_involved": sum(d.values()), "examples": [[tr(k, 60), n] for k, n in sorted(d.items())[:3]]}
P["duplicates"] = {"exact_rows": N - len(set(map(tuple, rows))),
    "docket": dups(col["Docket Number"]), "fr_docnum": dups([x or "" for x in fr]),
    "ad_number": dups(col["AD Number"]), "summary": dups(col["Summary"]),
    "subject_heading": dups(col["Subject Heading"])}
P["duplicates"]["exact_rows_excluding_docket"] = N - len(set(tuple(r[1:]) for r in rows))

# Distributions
def dist(h, k=15):
    c = Counter(x if x.strip() else "<empty>" for x in col[h])
    return {"distinct": len(c), "top": [[tr(a, 80), b] for a, b in sorted(c.items(), key=lambda kv: (-kv[1], kv[0]))[:k]]}
for h in ["Status", "Product Type", "Product Subtype", "CFR Part Reference", "Service/Office", "Office of Primary Responsibility", "Additional Reference", "AC Reference", "AIR Policy Reference"]:
    P.setdefault("dist", {})[h] = dist(h, 15)
P["dist"]["Make"] = dist("Make", 15)
P["dist"]["Type x Subtype"] = [[a, b] for a, b in sorted(Counter(f"{r[hdr.index('Product Type')]} / {r[hdr.index('Product Subtype')]}" for r in rows).items(), key=lambda kv: (-kv[1], kv[0]))]
# Models
mi = hdr.index("Model"); per = []; allm = Counter(); odd = {}
for r in rows:
    toks = [t.strip() for t in r[mi].split("|") if t.strip()]; per.append(len(toks)); allm.update(toks)
oddpats = {"comma": r",", "paren": r"[()]", "and_word": r"\band\b", "series": r"(?i)-series|series", "wildcard": r"(?i)\ball\b|\*|\bx{2,}\b|\bthrough\b|\bthru\b|\bvariants?\b", "slash": r"/", "space": r" "}
for k, p in oddpats.items():
    m = sorted(t for t in allm if re.search(p, t)); odd[k] = {"count": len(m), "examples": [tr(x, 60) for x in m[:5]]}
P["models"] = {"rows_empty": sum(1 for x in per if x == 0), "per_row_min": min(per), "per_row_median": statistics.median(per), "per_row_max": max(per),
    "per_row_total_tokens": sum(per), "distinct": len(allm), "odd_tokens": odd,
    "top5": [[a, b] for a, b in sorted(allm.items(), key=lambda kv: (-kv[1], kv[0]))[:5]],
    "empty_make_rows": sum(emp(x) for x in col["Make"])}
mk = hdr.index("Make"); P["models"]["make_with_pipe_rows"] = sum("|" in r[mk] for r in rows)

# Summary signals
S = col["Summary"]
kw = {"supersede": r"(?i)supersede", "revise": r"(?i)revise", "AD_yyyy-mm-dd": r"(?i)AD \d{4}-\d{2}-\d{2}", "withdraw": r"(?i)withdraw", "reopen": r"(?i)reopen", "supplemental": r"(?i)supplemental"}
P["summary_signals"] = {k: sum(bool(re.search(p, s)) for s in S) for k, p in kw.items()}
refs = [re.findall(r"\d{4}-\d{2}-\d{2}", s) for s in S]
P["summary_signals"]["rows_with_ad_ref"] = sum(bool(x) for x in refs)
P["summary_signals"]["total_ad_refs"] = sum(len(x) for x in refs)
P["summary_signals"]["distinct_ad_refs"] = len(set(a for x in refs for a in x))
P["summary_signals"]["ad_ref_examples"] = sorted(set(a for x in refs for a in x))[:3]
P["summary_signals"]["ad_refs_per_row_max"] = max(len(x) for x in refs)
P["summary_signals"]["first_sentence_openers"] = [[a, b] for a, b in Counter(" ".join(s.split()[:6]) for s in S).most_common(5)]
adn = col["AD Number"]; adne = [x for x in adn if not emp(x)]
P["ad_number"] = {"non_empty": len(adne), "fill_rate": round(len(adne) / N, 4), "format_ok_YYYY-MM-DD": sum(bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", x.strip())) for x in adne), "examples": adne[:3]}
app = {"serial number": r"(?i)serial number", "S/N": r"S/N", "MSN": r"\bMSN", "all ": r"(?i)\ball ", "certain": r"(?i)\bcertain\b", "except": r"(?i)\bexcept", "part number": r"(?i)part number", "P/N": r"P/N"}
P["applicability_signals"] = {k: sum(bool(re.search(p, s)) for s in S) for k, p in app.items()}
sn = [s for s in S if re.search(r"(?i)serial number|S/N|\bMSN", s)]
P["applicability_signals"]["serial_examples"] = [tr(s, 300) for s in sn[:3]]
SL = [len(s) for s in S if s.strip()]
term = lambda s: s.rstrip()[-1:] in ".!?\")"
P["summary_len"] = {"min": min(SL), "median": statistics.median(SL), "max": max(SL), "empty": N - len(SL),
    "no_terminal_punct": sum(not term(s) for s in S if s.strip()), "ends_with_ellipsis": sum(s.rstrip().endswith("...") for s in S),
    "hist": dict(sorted(Counter(min(l // 100 * 100, 1000) for l in SL).items())),
    "no_terminal_examples": [tr(s[-80:], 80) for s in S if s.strip() and not term(s)][:3],
    "len_300_count": sum(l == 300 for l in SL), "len_256_count": sum(l == 256 for l in SL)}

# Pages / Citation
pg = col["Page Number"]; tot = 0; fail = []; spans = []
for i, s in enumerate(pg):
    m = re.fullmatch(r"\s*\(Pages? (\d+)(?:-(\d+))?\)\s*", s)
    if not m: fail.append(tr(s, 40)); continue
    a = int(m.group(1)); b = int(m.group(2) or a); spans.append(b - a + 1)
    if b < a: fail.append("end<start:" + s)
P["pages"] = {"parsed": len(spans), "fail": len(fail), "fail_examples": fail[:5], "span_sum_inclusive": sum(spans), "span_min": min(spans), "span_max": max(spans), "span_median": statistics.median(spans),
    "single_page_rows": sum(1 for x in pg if re.search(r"\(Page \d+\)", x))}
ci = col["Citation"]
cf = Counter(re.sub(r"\d+", "N", re.sub(r"\((\w+), \w+ \d+, \d{4}\)", r"(\1, DATE)", s)) for s in ci)
P["citation"] = {"format_templates": cf.most_common(5), "regex_ok": sum(bool(re.fullmatch(r"Federal Register, Volume \d+ Number \d+ \(\w+day, \w+ \d+, \d{4}\)", s)) for s in ci)}
cm = 0; cc = 0
for s, d in zip(ci, col["Citation Publish Date"]):
    m = re.search(r"\((\w+), (\w+ \d+, \d{4})\)", s); dd = pdate(d)
    if m and dd:
        cc += 1
        try: cm += datetime.strptime(m.group(2), "%B %d, %Y") != dd
        except Exception: cm += 1
P["citation"]["date_vs_publish_compared"] = cc; P["citation"]["date_vs_publish_mismatch"] = cm
P["comments_col"] = {"non_empty": cols["Comments"]["non_empty"], "examples": [tr(x, 120) for x in col["Comments"] if x.strip()][:3]}

# Conflicts
iss = [pdate(x) for x in col["Issue Date"]]; pub = [pdate(x) for x in col["Citation Publish Date"]]; due = [pdate(x) for x in col["Comments Due Date"]]
cfl = {}
cfl["missing_docket"] = sum(emp(x) for x in col["Docket Number"])
cfl["missing_hyperlink_on_docket"] = sum("Docket Number" not in l for l in links)
cfl["missing_link_text"] = sum(emp(x) for x in urls)
cfl["issue_after_publish"] = sum(1 for a, b in zip(iss, pub) if a and b and a > b)
cfl["due_before_or_eq_publish"] = sum(1 for a, b in zip(pub, due) if a and b and b <= a)
dd = [(b - a).days for a, b in zip(pub, due) if a and b]
cfl["due_minus_publish_days"] = {"min": min(dd), "median": statistics.median(dd), "max": max(dd)} if dd else None
cfl["due_minus_publish_days_hist_top"] = Counter(dd).most_common(4)
cfl["docket_pattern_ok"] = sum(bool(re.fullmatch(r"FAA-\d{4}-\d+", x)) for x in col["Docket Number"])
cfl["status_by_due_vs_snapshot_max_publish"] = None
mx = max(x for x in pub if x)
cfl["status_vs_due_date"] = [[f"{s}|due<=max_publish:{(d is not None and d <= mx)}", n] for (s, d), n in sorted(Counter((s, d) for s, d in zip(col["Status"], due)).items(), key=lambda kv: (kv[0][0], str(kv[0][1])))][:0]
cfl["status_vs_due_past_max_publish"] = dict(Counter(f"{s}|due<=max_publish={d is not None and d <= mx}" for s, d in zip(col["Status"], due)))
cfl["subject_not_starting_airworthiness_directives"] = sum(not x.startswith("Airworthiness Directives") for x in col["Subject Heading"])
cfl["summary_not_starting_FAA_proposes"] = sum(not x.startswith("The FAA proposes") for x in S)
cfl["whitespace_padded_cells"] = sum(1 for r in rows for x in r if x != x.strip())
cfl["rows_sorted_by_publish_date_desc"] = all(a >= b for a, b in zip([p for p in pub if p], [p for p in pub if p][1:]))
P["conflicts"] = cfl

# ---- Population split (modern = Service/Office filled) ----
newm = [not emp(x) for x in col["Service/Office"]]
def cnt(f): return {"modern": sum(1 for i in range(N) if newm[i] and f(i)), "legacy": sum(1 for i in range(N) if not newm[i] and f(i))}
POP = {"n": {"modern": sum(newm), "legacy": N - sum(newm)}}
def pr(h): return [col[h][i] for i in range(N)]
for h in ["Model", "AD Number", "Product Type", "Product Subtype", "Office of Primary Responsibility", "Comments Due Date", "Summary", "Comments", "Make"]:
    POP["nonempty_" + h] = cnt(lambda i, h=h: not emp(col[h][i]))
yr = lambda s: s[-4:]
for nm, flag in (("modern", True), ("legacy", False)):
    ps = [pdate(col["Citation Publish Date"][i]) for i in range(N) if newm[i] == flag]
    POP["publish_range_" + nm] = [min(ps).strftime("%Y-%m-%d"), max(ps).strftime("%Y-%m-%d")]
    POP["publish_year_counts_" + nm] = dict(sorted(Counter(str(d.year) for d in ps).items()))
    POP["summary_len_median_" + nm] = statistics.median([len(col["Summary"][i]) for i in range(N) if newm[i] == flag and col["Summary"][i].strip()])
def lk(u):
    if emp(u): return "empty"
    if u.startswith("https://www.federalregister.gov/documents/"): return "fr_documents_url"
    if re.fullmatch(r"https?://\S+/d/\d{4}-\d+", u.strip()): return "fr_short_url_/d/"
    if re.search(r"^\"?Withdrawal: ?\"?https?://", u): return "withdrawal_prefix_plus_short_url"
    if u.startswith("http"): return "other_url"
    return "free_text_nonurl"
LK = Counter((("modern" if newm[i] else "legacy"), lk(urls[i])) for i in range(N))
POP["link_text_kinds"] = {f"{a}:{b}": n for (a, b), n in sorted(LK.items())}
POP["link_text_free_text_len"] = {"min": min(len(u) for u in urls if lk(u) == "free_text_nonurl"), "median": statistics.median([len(u) for u in urls if lk(u) == "free_text_nonurl"]), "max": max(len(u) for u in urls if lk(u) == "free_text_nonurl")}
POP["link_text_free_text_starts"] = [[tr(a.replace("\n", " ").replace("\t", " "), 40), b] for a, b in Counter(u[:25] for u in urls if lk(u) == "free_text_nonurl").most_common(4)]
POP["link_text_literal_backslash_n_rows"] = sum("\\n" in u for u in urls)
POP["summary_literal_backslash_n_rows"] = sum("\\n" in x for x in S)
POP["summary_real_newline_rows"] = sum("\n" in x for x in S)
shortdoc = []
for u in urls:
    m = re.search(r"federalregister\.gov/d/(\d{4}-\d+)", u)
    shortdoc.append(m.group(1) if m else None)
POP["fr_docnum_from_short_url"] = {"rows": sum(bool(x) for x in shortdoc), "unique": len(set(x for x in shortdoc if x)), "overlap_with_documents_url_docnums": len(set(x for x in shortdoc if x) & set(frs))}
alld = [a or b for a, b in zip(fr, shortdoc)]
POP["fr_docnum_any_source"] = {"rows": sum(bool(x) for x in alld), "unique": len(set(x for x in alld if x))}
POP["drs_target_kinds"] = dict(Counter("FR-ADNPRM-YYYY-N" if re.search(r"FR-ADNPRM-\d{4}-\d+-\d+\.", t) else ("FR-ADNPRM-R1-YYYY-N" if "FR-ADNPRM-R" in t else "other_hex_id") for t in tg))
POP["drs_target_len"] = {"min": min(map(len, tg)), "max": max(map(len, tg))}
POP["drs_target_unique_per_row"] = len(set(tg)) == N
mmx = []
for i, (t, f) in enumerate(zip(tg, fr)):
    m = re.search(r"FR-ADNPRM-(\d{4}-\d+)-", t or "")
    if m and f and m.group(1) != f: mmx.append([m.group(1), f, col["Docket Number"][i]])
POP["drs_vs_url_docnum_mismatch_rows"] = mmx
# page formats
POP["page_formats"] = dict(Counter(re.sub(r"\d+", "N", x) for x in pg))
rel = []; badp = []
for x in pg:
    m = re.search(r"(\d+)(?:-(\d+))?", x)
    a = int(m.group(1)); b = int(m.group(2) or a); (rel.append(b - a + 1) if b >= a else badp.append(x))
POP["page_span_relaxed"] = {"parsed": len(rel), "end_lt_start": badp[:5], "sum_inclusive": sum(rel), "median": statistics.median(rel), "max": max(rel)}
POP["page_span_sum_modern_only"] = None
ms = []
for i, x in enumerate(pg):
    m = re.search(r"(\d+)(?:-(\d+))?", x)
    if newm[i]: ms.append(int(m.group(2) or m.group(1)) - int(m.group(1)) + 1)
POP["page_span_sum_modern_only"] = {"sum": sum(ms), "rows": len(ms)}
POP["citation_formats"] = {"modern_FR_Volume_first": sum(s.startswith("Federal Register, Volume") for s in ci), "legacy_FR_colon_date_first": sum(s.startswith("Federal Register:") for s in ci), "other": sum(not s.startswith("Federal Register") for s in ci),
    "legacy_example": [x for x in ci if x.startswith("Federal Register:")][:1]}
cm2 = 0; cc2 = 0; ex = []
for s, d in zip(ci, col["Citation Publish Date"]):
    m = re.search(r"((?:January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}, \d{4})", s); dd = pdate(d)
    if m and dd:
        cc2 += 1
        if datetime.strptime(m.group(1), "%B %d, %Y") != dd:
            cm2 += 1; ex.append([m.group(1), d]) if len(ex) < 3 else 0
POP["citation_date_vs_publish_relaxed"] = {"compared": cc2, "mismatch": cm2, "examples": ex}
POP["status_Current_but_Summary_withdraw_rows"] = sum(bool(re.search(r"(?i)withdraw", x)) for x in S)
POP["withdraw_rows_by_pop"] = cnt(lambda i: bool(re.search(r"(?i)withdraw", S[i])))
POP["supersede_rows_by_pop"] = cnt(lambda i: bool(re.search(r"(?i)supersede", S[i])))
POP["serial_rows_by_pop"] = cnt(lambda i: bool(re.search(r"(?i)serial number|S/N|\bMSN", S[i])))
POP["comments_due_before_publish_by_pop"] = cnt(lambda i: bool(pub[i] and due[i] and due[i] < pub[i]))
POP["comments_due_empty_by_pop"] = cnt(lambda i: due[i] is None)
POP["docket_types_by_pop"] = {"modern": dict(Counter(re.sub(r"\d", "N", col["Docket Number"][i]) for i in range(N) if newm[i]).most_common(4)), "legacy": dict(Counter(re.sub(r"\d", "N", col["Docket Number"][i]) for i in range(N) if not newm[i]).most_common(5))}
POP["docket_dup_by_pop"] = {"dup_rows_modern": sum(1 for i in range(N) if newm[i] and Counter(col["Docket Number"])[col["Docket Number"][i]] > 1), "dup_rows_legacy": sum(1 for i in range(N) if not newm[i] and Counter(col["Docket Number"])[col["Docket Number"][i]] > 1)}
POP["comments_all"] = [x for x in col["Comments"] if x.strip()]
POP["multi_value_product_rows"] = {"type_pipe": sum("|" in x for x in col["Product Type"]), "make_pipe": sum("|" in x for x in col["Make"])}
POP["models_modern_rows_with_model"] = sum(1 for i in range(N) if newm[i] and not emp(col["Model"][i]))
POP["model_empty_modern_nonwithdraw"] = sum(1 for i in range(N) if newm[i] and emp(col["Model"][i]) and not re.search(r"(?i)withdraw", S[i]))
P["population"] = POP
json.dump(P, open(out, "w"), indent=1, sort_keys=True, default=str, ensure_ascii=False)
